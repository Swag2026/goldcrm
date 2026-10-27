"""Run: GERP_MOCK=1 DATABASE_URL=sqlite:///./test.db ADMIN_PASSWORD=admin12345 python3 -m pytest -q test_api.py
(ya seedha: python3 test_api.py)"""
import io
import os
import tempfile

_tmp = tempfile.mkdtemp()
os.environ.update({
    "GERP_MOCK": "1", "DATABASE_URL": f"sqlite:///{_tmp}/t.db", "ADMIN_USERNAME": "admin",
    "ADMIN_PASSWORD": "admin12345", "SECRET_KEY": "test-secret-key-test-secret-key-1234",
})

from datetime import timedelta  # noqa: E402

from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

import config  # noqa: E402
from main import app  # noqa: E402


def H(tok):
    return {"Authorization": f"Bearer {tok}"}


def run():
    with TestClient(app) as c:
        # --- auth
        assert c.get("/api/customers").status_code == 401
        assert c.post("/api/auth/login", json={"username": "admin", "password": "bad"}).status_code == 401
        r = c.post("/api/auth/login", json={"username": "admin", "password": "admin12345"})
        assert r.status_code == 200, r.text
        adm = r.json()["token"]
        assert c.get("/api/auth/me", headers=H(adm)).json()["role"] == "admin"

        # --- users
        r = c.post("/api/users", headers=H(adm), json={"username": "Ali", "name": "Ali K", "password": "ali-pass-123"})
        assert r.status_code == 200, r.text
        ali_id = r.json()["id"]
        assert c.post("/api/users", headers=H(adm), json={"username": "ali", "name": "x", "password": "ali-pass-123"}).status_code == 409
        ali = c.post("/api/auth/login", json={"username": "ali", "password": "ali-pass-123"}).json()["token"]
        assert c.get("/api/users", headers=H(ali)).status_code == 403
        assert len(c.get("/api/collectors", headers=H(ali)).json()) == 2
        assert c.patch("/api/users/1", headers=H(adm), json={"active": False}).status_code == 400  # self-lock

        # --- customers + sync (mock = 40 customers)
        r = c.get("/api/customers?flt=all&limit=100", headers=H(adm)).json()
        assert r["total"] == 40, r["total"]
        rows_bal = c.get("/api/customers", headers=H(adm)).json()["items"]
        assert 0 < len(rows_bal) < 40
        golds = [x["gold_balance"] for x in rows_bal]
        assert golds == sorted(golds, reverse=True)
        gd = c.get("/api/customers?flt=gold_debtors", headers=H(adm)).json()["items"]
        assert gd and all(x["gold_balance"] > 0.001 for x in gd)
        owe = c.get("/api/customers?flt=we_owe", headers=H(adm)).json()["items"]
        assert all(min(x["gold_balance"], x["cash_balance"], x["silver_balance"]) < 0 for x in owe)
        q = c.get("/api/customers?flt=all&q=" + "الركن", headers=H(adm)).json()["items"]
        assert q and all("الركن" in x["name"] for x in q)
        assert c.get("/api/customers?flt=all&q=purest", headers=H(adm)).json()["total"] > 0  # English bracket text
        assert c.get("/api/customers?flt=all&q=11210001", headers=H(adm)).json()["total"] == 1
        assert c.post("/api/sync", headers=H(adm)).json()["synced"] is False  # < 30s throttle

        target = gd[0]["id"]

        # --- assign permissions
        assert c.patch(f"/api/customers/{target}", headers=H(ali), json={"assigned_to": 1}).status_code == 403
        assert c.patch(f"/api/customers/{target}", headers=H(ali), json={"assigned_to": ali_id}).status_code == 200
        assert c.get("/api/customers?assigned=me", headers=H(ali)).json()["total"] == 1
        assert c.patch(f"/api/customers/{target}", headers=H(ali), json={"status": "bogus"}).status_code == 400
        assert c.patch(f"/api/customers/{target}", headers=H(ali), json={"status": "hold", "notes": "n1"}).status_code == 200
        assert c.post("/api/customers/assign", headers=H(ali), json={"ids": [target], "user_id": 1}).status_code == 403
        assert c.post("/api/customers/assign", headers=H(adm), json={"ids": [gd[1]["id"], gd[2]["id"]], "user_id": ali_id}).json()["updated"] == 2

        # --- followups & promises
        t = config.today()
        bad = c.post(f"/api/customers/{target}/followups", headers=H(ali), json={"kind": "call", "promise_gold": 5})
        assert bad.status_code == 400
        bad = c.post(f"/api/customers/{target}/followups", headers=H(ali),
                     json={"kind": "call", "promise_date": t.isoformat()})
        assert bad.status_code == 400
        assert c.post(f"/api/customers/{target}/followups", headers=H(ali), json={"kind": "xx"}).status_code == 400
        f1 = c.post(f"/api/customers/{target}/followups", headers=H(ali), json={
            "kind": "call", "outcome": "promised", "note": "wada kiya",
            "next_followup_date": (t - timedelta(days=2)).isoformat(),
            "promise_date": (t - timedelta(days=1)).isoformat(), "promise_gold": 100.5, "promise_cash": 5000})
        assert f1.status_code == 200, f1.text
        assert f1.json()["promise_status"] == "pending"
        fid = f1.json()["id"]

        d = c.get(f"/api/customers/{target}", headers=H(ali)).json()
        assert d["status"] == "hold" and d["notes"] == "n1"
        assert d["promise"]["overdue"] is True and d["promise"]["gold"] == 100.5
        assert len(d["followups"]) == 1 and d["followups"][0]["user_name"] == "Ali K"
        assert "gold_debit" in d

        md = c.get("/api/myday", headers=H(ali)).json()
        assert md["scope"] == "mine"
        assert any(x["id"] == target and x["days_overdue"] == 2 for x in md["followups"])
        assert any(x["id"] == target and x["days_overdue"] == 1 for x in md["promises"])
        assert c.get("/api/myday", headers=H(adm)).json()["scope"] == "all"

        # --- promise resolution
        assert c.patch(f"/api/followups/{fid}", headers=H(ali), json={"promise_status": "weird"}).status_code == 400
        assert c.patch(f"/api/followups/{fid}", headers=H(ali), json={"promise_status": "broken"}).status_code == 200
        md = c.get("/api/myday", headers=H(ali)).json()
        assert not any(x["id"] == target for x in md["promises"])
        # naya followup next_followup future => My Day se hat jata hai
        c.post(f"/api/customers/{target}/followups", headers=H(ali), json={
            "kind": "whatsapp", "outcome": "reached", "next_followup_date": (t + timedelta(days=5)).isoformat()})
        assert not any(x["id"] == target for x in c.get("/api/myday", headers=H(ali)).json()["followups"])

        # --- delete rules
        other = c.post("/api/users", headers=H(adm), json={"username": "bob", "name": "Bob", "password": "bob-pass-123"}).json()
        bob = c.post("/api/auth/login", json={"username": "bob", "password": "bob-pass-123"}).json()["token"]
        assert c.delete(f"/api/followups/{fid}", headers=H(bob)).status_code == 403
        assert c.delete(f"/api/followups/{fid}", headers=H(ali)).status_code == 200

        # --- overview
        ov = c.get("/api/overview", headers=H(adm)).json()
        tt = ov["totals"]
        db_rows = c.get("/api/customers?flt=all&limit=100", headers=H(adm)).json()["items"]
        exp_gold = round(sum(x["gold_balance"] for x in db_rows if x["gold_balance"] > 0.001), 4)
        assert abs(tt["gold_receivable"] - exp_gold) < 0.01, (tt["gold_receivable"], exp_gold)
        assert ov["customers_total"] == 40 and ov["top_gold"][0]["value"] >= ov["top_gold"][-1]["value"]
        assert len(ov["activity"]) == 14 and len(ov["trend"]) == 1
        assert ov["promises_broken_30d"] == 1 or ov["promises_broken_30d"] == 0  # f1 delete ho chuka

        # --- kpi (mock rows => Colab wale numbers)
        k = c.get(f"/api/kpi?from_date=2026-01-01&to_date={t.isoformat()}", headers=H(adm)).json()
        assert abs(k["gold_grams"]["current"] - 11570.06) < 0.01
        assert abs(k["gold_cash"]["current"] - 7027366.03) < 0.01
        assert abs(k["silver"]["current"] - 557000.0) < 0.01
        assert "331110001" in k["missing_codes"]
        assert c.get("/api/kpi?from_date=bad&to_date=2026-01-01", headers=H(adm)).status_code == 422

        # --- statement
        sq = f"from_date={(t - timedelta(days=30)).isoformat()}&to_date={t.isoformat()}"
        st = c.get(f"/api/customers/{target}/statement?{sq}", headers=H(adm)).json()
        for mk, v in st["metals"].items():
            assert abs(v["opening"] + v["debit"] - v["credit"] - v["closing"]) < 0.01, (mk, v)
        assert st["metals"]["gold"]["closing"] > 0 and st["customer"]["code"]
        assert c.get(f"/api/customers/{target}/statement?from_date={t.isoformat()}&to_date={(t - timedelta(days=1)).isoformat()}", headers=H(adm)).status_code == 400
        assert c.get(f"/api/customers/999999/statement?{sq}", headers=H(adm)).status_code == 404
        assert c.get(f"/api/customers/{target}/statement?{sq}").status_code == 401
        rx = c.get(f"/api/customers/{target}/statement.xlsx?{sq}", headers=H(adm))
        assert rx.status_code == 200 and load_workbook(io.BytesIO(rx.content)).active.max_row >= 9

        # --- export
        r = c.get("/api/customers/export.xlsx?flt=gold_debtors", headers=H(adm))
        assert r.status_code == 200 and r.headers["content-type"].startswith("application/vnd.openxml")
        wb = load_workbook(io.BytesIO(r.content))
        assert wb.active.max_row == len(gd) + 1

        # --- password change
        assert c.post("/api/auth/password", headers=H(ali), json={"current": "nope", "new": "newpass-1234"}).status_code == 400
        assert c.post("/api/auth/password", headers=H(ali), json={"current": "ali-pass-123", "new": "newpass-1234"}).status_code == 200
        assert c.post("/api/auth/login", json={"username": "ali", "password": "newpass-1234"}).status_code == 200

        # --- disabled user lock-out
        c.patch(f"/api/users/{other['id']}", headers=H(adm), json={"active": False})
        assert c.get("/api/auth/me", headers=H(bob)).status_code == 401

        # --- history: pehle din koi delta nahi
        lst = c.get("/api/customers?flt=all&limit=100", headers=H(adm)).json()
        assert lst["prev_day"] is None and lst["items"][0]["delta"] is None
        d0 = c.get(f"/api/customers/{target}", headers=H(adm)).json()
        assert len(d0["history"]) == 1 and d0["history"][0]["day"] == config.today().isoformat()
        ov0 = c.get("/api/overview", headers=H(adm)).json()
        assert ov0["changes"] is None and ov0["paid_down"]["gold"] == 0 and ov0["prev_day"] is None

        # --- kal ka snapshot banao: target ka gold kal 1.5x tha => aaj 'paid down'
        from sqlalchemy import select
        from db import CompanySnapshot, Customer, CustomerSnapshot, SessionLocal
        y = config.today() - timedelta(days=1)
        with SessionLocal() as db:
            tb = db.get(Customer, target).gold_balance
            gold_rec = 0.0
            for cu in db.scalars(select(Customer)):
                m = 1.5 if cu.id == target else 1.0
                db.add(CustomerSnapshot(customer_id=cu.id, day=y, cash_balance=cu.cash_balance * m,
                                        gold_balance=cu.gold_balance * m, silver_balance=cu.silver_balance * m))
                gold_rec += max(cu.gold_balance, 0)
            db.add(CompanySnapshot(company="swag", day=y, gold_receivable=gold_rec - 100, debtors=0))
            ysnap_total = gold_rec
            db.commit()
        assert tb > 0
        ov1 = c.get("/api/overview", headers=H(adm)).json()
        assert ov1["prev_day"] == y.isoformat()
        assert abs(ov1["paid_down"]["gold"] - 0.5 * tb) < 0.01, (ov1["paid_down"]["gold"], 0.5 * tb)
        assert ov1["movers"]["gold"]["down"][0]["id"] == target
        assert abs(ov1["changes"]["gold_receivable"] - 100) < 0.01
        lst = c.get("/api/customers?flt=all&limit=100", headers=H(adm)).json()
        row = next(x for x in lst["items"] if x["id"] == target)
        assert abs(row["delta"]["gold"] - (-0.5 * tb)) < 0.01 and lst["prev_day"] == y.isoformat()
        assert len(c.get(f"/api/customers/{target}", headers=H(adm)).json()["history"]) == 2
        board = {b["name"]: b for b in ov1["board"]}
        assert "Ali K" in board and board["Ali K"]["followups_7d"] >= 1

        # --- customer GERP list se gayab ho jaye => balance zero (totals stale na rahein), history safe
        import gerp
        import services
        orig = gerp.mock_customers
        gerp.mock_customers = lambda d=None, i=0: [x for x in orig(d, i) if x["id"] + (i * 1_000_000) != target]
        try:
            with SessionLocal() as db:
                services.sync_customers(db, "swag", force=True)
                cu = db.get(Customer, target)
                assert cu.gold_balance == 0 and cu.cash_balance == 0 and cu.silver_balance == 0
                snap = db.get(CustomerSnapshot, (target, config.today()))
                assert snap.gold_balance == 0
        finally:
            gerp.mock_customers = orig

    multi_company_tests()
    migration_test()
    print("ALL BACKEND TESTS PASSED")


def multi_company_tests():
    from sqlalchemy import select as sel
    with TestClient(app) as c:
        adm = c.post("/api/auth/login", json={"username": "admin", "password": "admin12345"}).json()["token"]
        comps = c.get("/api/companies", headers=H(adm)).json()
        assert [x["key"] for x in comps] == ["swag", "laroche"] and all(set(x) == {"key", "name"} for x in comps)
        assert c.get("/api/companies").status_code == 401
        sw = c.get("/api/customers?flt=all&limit=100&company=swag", headers=H(adm)).json()
        la = c.get("/api/customers?flt=all&limit=100&company=laroche", headers=H(adm)).json()
        assert sw["total"] == 40 and la["total"] == 40
        sw_ids = {x["id"] for x in sw["items"]}
        la_ids = {x["id"] for x in la["items"]}
        assert not (sw_ids & la_ids), "ids overlap across companies"
        assert all(i < 1_000_000 for i in sw_ids) and all(1_000_000 <= i < 2_000_000 for i in la_ids)
        assert all(x["company"] == "laroche" for x in la["items"])
        assert sw["sync"]["company"] == "swag" and la["sync"]["company"] == "laroche"
        assert c.get("/api/customers?company=nope", headers=H(adm)).status_code == 404
        # totals alag
        ov_s = c.get("/api/overview?company=swag", headers=H(adm)).json()
        ov_l = c.get("/api/overview?company=laroche", headers=H(adm)).json()
        assert ov_s["company"] == "swag" and ov_l["company"] == "laroche"
        assert abs(ov_s["totals"]["gold_receivable"] - ov_l["totals"]["gold_receivable"]) > 1
        assert ov_l["customers_total"] == 40 and ov_l["feed"] == [] and ov_l["promises_pending"] == 0, "swag activity leaked into laroche"
        assert len(ov_l["trend"]) == 1
        # followup on laroche customer stays in laroche
        lid = next(x["id"] for x in la["items"] if x["gold_balance"] > 1)
        assert c.post(f"/api/customers/{lid}/followups", headers=H(adm), json={"kind": "call", "outcome": "reached", "note": "x"}).status_code == 200
        ov_l = c.get("/api/overview?company=laroche", headers=H(adm)).json()
        ov_s = c.get("/api/overview?company=swag", headers=H(adm)).json()
        assert len(ov_l["feed"]) == 1 and all(f["customer_id"] != lid for f in ov_s["feed"])
        # customer detail + statement use their own company
        d = c.get(f"/api/customers/{lid}", headers=H(adm)).json()
        assert d["company"] == "laroche"
        t0 = config.today()
        st = c.get(f"/api/customers/{lid}/statement?from_date={(t0 - timedelta(days=30)).isoformat()}&to_date={t0.isoformat()}", headers=H(adm)).json()
        assert st["company"] == "LAROUCHE" and abs(st["metals"]["gold"]["closing"] - d["gold_balance"]) < 0.01
        # KPI per company
        ks = c.get(f"/api/kpi?from_date=2026-01-01&to_date={t0.isoformat()}&company=swag", headers=H(adm)).json()
        kl = c.get(f"/api/kpi?from_date=2026-01-01&to_date={t0.isoformat()}&company=laroche", headers=H(adm)).json()
        assert ks["company"] == "swag" and kl["company"] == "laroche"
        assert abs(kl["gold_grams"]["current"] - 11570.06 * 1.6) < 0.05 and abs(ks["gold_grams"]["current"] - 11570.06) < 0.05
        # sync status/ttl independent
        assert c.post("/api/sync?company=laroche", headers=H(adm)).json()["synced"] is False  # <30s throttle
        # export
        rx = c.get("/api/customers/export.xlsx?flt=all&company=laroche", headers=H(adm))
        assert rx.status_code == 200 and "laroche" in rx.headers["content-disposition"]
        assert load_workbook(io.BytesIO(rx.content)).active.max_row == 41
        # my day scoped
        md = c.get("/api/myday?scope=all&company=swag", headers=H(adm)).json()
        assert all(x["id"] < 1_000_000 for x in md["followups"] + md["promises"] + md["unassigned_debtors"])


def migration_test():
    """Purana single-company DB (company column + company_snapshots nahi) => data safe, naya schema."""
    import subprocess, sqlite3, sys
    path = os.path.join(_tmp, "old.db")
    cx = sqlite3.connect(path)
    cx.executescript("""
    CREATE TABLE users (id INTEGER PRIMARY KEY, username VARCHAR(64) UNIQUE, name VARCHAR(120), password_hash VARCHAR(255), role VARCHAR(16), active BOOLEAN, created_at DATETIME);
    CREATE TABLE customers (id INTEGER PRIMARY KEY, code VARCHAR(64), name VARCHAR(255), name_ar VARCHAR(255), account_id INTEGER,
      cash_debit FLOAT, cash_credit FLOAT, cash_balance FLOAT, gold_debit FLOAT, gold_credit FLOAT, gold_balance FLOAT,
      silver_debit FLOAT, silver_credit FLOAT, silver_balance FLOAT, synced_at DATETIME, assigned_to INTEGER, status VARCHAR(16), notes TEXT);
    CREATE TABLE snapshots (day DATE PRIMARY KEY, gold_receivable FLOAT, gold_payable FLOAT, cash_receivable FLOAT, cash_payable FLOAT,
      silver_receivable FLOAT, silver_payable FLOAT, debtors INTEGER);
    INSERT INTO customers (id, code, name, name_ar, gold_balance, status, notes) VALUES (17, '11210001', 'Old Customer', 'عميل قديم', 12.5, 'hold', 'keep me');
    INSERT INTO snapshots VALUES ('2026-10-01', 100, 0, 5000, 0, 0, 0, 3);
    """)
    cx.commit(); cx.close()
    code = (
        "import os, sqlite3\n"
        f"os.environ['DATABASE_URL']='sqlite:///{path}'; os.environ['GERP_MOCK']='1'\n"
        "import db; db.init_db(); db.init_db()\n"  # do baar: idempotent
        f"cx = sqlite3.connect('{path}')\n"
        "row = cx.execute('select company, name, status, notes, gold_balance from customers where id=17').fetchone()\n"
        "assert row == ('swag', 'Old Customer', 'hold', 'keep me', 12.5), row\n"
        "snaps = cx.execute('select company, day, gold_receivable, debtors from company_snapshots').fetchall()\n"
        "assert snaps == [('swag', '2026-10-01', 100.0, 3)], snaps\n"
        "print('migration ok')\n"
    )
    r = subprocess.run([sys.executable, "-c", code], capture_output=True, text=True, cwd=os.path.dirname(os.path.abspath(__file__)))
    assert r.returncode == 0 and "migration ok" in r.stdout, (r.stdout, r.stderr[-800:])


def test_all():
    run()


if __name__ == "__main__":
    run()
