"""GERP 'Customer and Supplier Report' (Excel export) import: address, mobile, last invoice, last transaction date.

API mein ye fields nahi aate, isliye admin report export karke upload karta hai. Match: account code == customer code.
"""
import io
import re
from datetime import date, datetime
from typing import Optional

from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

import config
from db import Customer

HEADERS = {
    "account code": "code", "account": "name", "address": "address", "mobile no": "mobile", "mobile": "mobile",
    "last invoice amt": "last_invoice_amt", "last transaction date": "last_txn_date",
    "balance cash": "cash", "balance gold": "gold", "balance silver": "silver",
}
MAX_BYTES = 8 * 1024 * 1024


def norm_name(s) -> str:
    s = re.sub(r"\s+", " ", str(s or "")).strip().lower()
    return re.sub(r"\s*([()])\s*", r"\1", s)


def _num(v) -> Optional[float]:
    if v is None or v == "":
        return None
    try:
        return float(str(v).replace(",", ""))
    except ValueError:
        return None


def _date(v) -> Optional[date]:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v or "").strip()[:10]
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def _mobile(v) -> str:
    if v is None or v == "":
        return ""
    if isinstance(v, (int, float)):  # Excel number => leading zero kho jata hai (5xxxxxxxx)
        v = str(int(v))
        if len(v) == 9 and v.startswith("5"):
            v = "0" + v
    return re.sub(r"[^\d+]", "", str(v))[:64]


def parse_report(data: bytes) -> dict:
    if len(data) > MAX_BYTES:
        raise HTTPException(413, "File bahut badi hai (max 8 MB)")
    try:
        ws = load_workbook(io.BytesIO(data), data_only=True, read_only=True).worksheets[0]
    except Exception:  # noqa: BLE001
        raise HTTPException(400, "Ye valid .xlsx file nahi lagti")
    report_date, cols, rows = None, None, []
    for r in ws.iter_rows(values_only=True):
        cells = list(r)
        first = str(cells[0] or "").strip() if cells else ""
        if cols is None:
            m = re.match(r"Date Until\s*:\s*(.+)", first, re.I)
            if m:
                report_date = _date(m.group(1).strip())
            labels = [str(c or "").strip().lower() for c in cells]
            if "account code" in labels:
                cols = {HEADERS[l]: i for i, l in enumerate(labels) if l in HEADERS}
            continue
        code = str(cells[cols["code"]] or "").strip() if "code" in cols and cols["code"] < len(cells) else ""
        if not code or not re.fullmatch(r"\d{4,}", code):
            continue  # totals / khali rows

        def cell(k):
            i = cols.get(k)
            return cells[i] if i is not None and i < len(cells) else None
        rows.append({
            "code": code, "name": str(cell("name") or "").strip(), "address": re.sub(r"\s+", " ", str(cell("address") or "")).strip()[:500],
            "mobile": _mobile(cell("mobile")), "last_invoice_amt": _num(cell("last_invoice_amt")),
            "last_txn_date": _date(cell("last_txn_date")),
            "cash": _num(cell("cash")) or 0.0, "gold": _num(cell("gold")) or 0.0, "silver": _num(cell("silver")) or 0.0,
        })
    if cols is None or "code" not in cols:
        raise HTTPException(400, "Header row 'Account Code' nahi mili: GERP ki 'Customer and Supplier Report' ka Excel upload karo")
    return {"report_date": report_date, "rows": rows}


def apply_report(db: Session, company: str, parsed: dict, force: bool = False) -> dict:
    customers = {c.code: c for c in db.scalars(select(Customer).where(Customer.company == company))}
    matched, same_name, unmatched = [], 0, []
    for r in parsed["rows"]:
        c = customers.get(r["code"])
        if c is None:
            unmatched.append(r)
            continue
        matched.append((c, r))
        if norm_name(c.name) == norm_name(r["name"]) or norm_name(c.name_ar) == norm_name(r["name"]):
            same_name += 1
    name_pct = (same_name / len(matched) * 100) if matched else 0.0
    if not force and len(matched) >= 5 and name_pct < 50:
        raise HTTPException(409, f"Ye file is company ({company}) ki nahi lagti: {len(matched)} codes mile par sirf {name_pct:.0f}% naam match hue. "
                                 "Sahi company select karke dobara upload karo.")
    if not matched:
        raise HTTPException(409, f"Is file ka koi account code '{company}' ke customers se match nahi hua. Sahi company select hai? Pehle ek baar sync karo.")
    now = config.now()
    for c, r in matched:
        c.address, c.mobile = r["address"], r["mobile"]
        c.last_invoice_amt, c.last_txn_date, c.profile_updated_at = r["last_invoice_amt"], r["last_txn_date"], now
    db.commit()
    return {
        "rows": len(parsed["rows"]), "matched": len(matched), "name_match_pct": round(name_pct, 1),
        "with_mobile": sum(1 for _, r in matched if r["mobile"]),
        "unmatched_total": len(unmatched),
        "unmatched": [{"code": r["code"], "name": r["name"]} for r in unmatched[:12]],
        "report_date": parsed["report_date"].isoformat() if parsed["report_date"] else None,
    }
