"""Gold Collections API — GERP balances + collections workflow."""
import io
import logging
import secrets
from datetime import date, datetime, timedelta
from typing import List, Optional

from fastapi import Depends, FastAPI, HTTPException, Query, Request
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from pydantic import BaseModel, Field
from sqlalchemy import select
from sqlalchemy.orm import Session

import config
import gerp
import report_import
import services as svc
from db import Customer, Followup, SessionLocal, User, get_db, init_db
from security import (admin_only, create_token, current_user, hash_password, throttle_check, throttle_clear,
                      throttle_fail, verify_password)

log = logging.getLogger("gold")
logging.basicConfig(level=logging.INFO)

KINDS = {"call", "whatsapp", "visit", "note", "other"}
OUTCOMES = {"reached", "no_answer", "promised", "refused", "paid", "other"}
STATUSES = {"active", "hold", "legal"}
PROMISE_STATUSES = {"pending", "kept", "broken"}

app = FastAPI(title="Gold Collections API")
app.add_middleware(CORSMiddleware, allow_origins=config.CORS_ORIGINS, allow_methods=["*"], allow_headers=["*"])


@app.on_event("startup")
def startup():
    init_db()
    if config.SECRET_KEY_IS_RANDOM:
        log.warning("SECRET_KEY set nahi hai: random key use ho rahi hai, restart par sab logout ho jayenge")
    with SessionLocal() as db:
        if not db.scalar(select(User).limit(1)):
            pw = config.ADMIN_PASSWORD or secrets.token_urlsafe(9)
            db.add(User(username=config.ADMIN_USERNAME, name="Admin", role="admin", password_hash=hash_password(pw)))
            db.commit()
            if config.ADMIN_PASSWORD:
                log.info("Admin user '%s' bana diya gaya", config.ADMIN_USERNAME)
            else:
                log.warning("Admin user '%s' ka temporary password: %s  (login ke baad badal lo)", config.ADMIN_USERNAME, pw)


def user_out(u: User) -> dict:
    return {"id": u.id, "username": u.username, "name": u.name, "role": u.role, "active": u.active}


def ensure_fresh(db: Session, company: str):
    """Cache purana ho to GERP se refresh (fail ho to purana cache hi chalta hai)."""
    svc.sync_customers(db, company, force=False, raise_errors=False)


def resolve_company(company: str = "") -> str:
    return svc.get_co(company)["key"]


@app.get("/api/health")
def health():
    return {"ok": True, "mock": config.MOCK}


@app.get("/api/companies")
def companies(_: User = Depends(current_user)):
    """Password kabhi bahar nahi jaata."""
    return [{"key": c["key"], "name": c["name"]} for c in config.COMPANIES]


# ------------------------------------------------------------------ auth
class LoginIn(BaseModel):
    username: str
    password: str


@app.post("/api/auth/login")
def login(body: LoginIn, request: Request, db: Session = Depends(get_db)):
    key = f"{request.client.host if request.client else '-'}|{body.username.lower()}"
    throttle_check(key)
    u = db.scalar(select(User).where(User.username == body.username.strip().lower()))
    if not u or not u.active or not verify_password(body.password, u.password_hash):
        throttle_fail(key)
        raise HTTPException(401, "Invalid username or password")
    throttle_clear(key)
    return {"token": create_token(u), "user": user_out(u)}


@app.get("/api/auth/me")
def me(user: User = Depends(current_user)):
    return user_out(user)


class PasswordIn(BaseModel):
    current: str
    new: str = Field(min_length=8)


@app.post("/api/auth/password")
def change_password(body: PasswordIn, user: User = Depends(current_user), db: Session = Depends(get_db)):
    if not verify_password(body.current, user.password_hash):
        raise HTTPException(400, "Current password is wrong")
    user.password_hash = hash_password(body.new)
    db.commit()
    return {"ok": True}


# ------------------------------------------------------------------ users (admin)
class UserIn(BaseModel):
    username: str = Field(min_length=3, max_length=64)
    name: str = Field(min_length=1, max_length=120)
    password: str = Field(min_length=8)
    role: str = "collector"


class UserPatch(BaseModel):
    name: Optional[str] = None
    role: Optional[str] = None
    active: Optional[bool] = None
    password: Optional[str] = Field(default=None, min_length=8)


@app.get("/api/users")
def list_users(_: User = Depends(admin_only), db: Session = Depends(get_db)):
    return [user_out(u) for u in db.scalars(select(User).order_by(User.id))]


@app.post("/api/users")
def create_user(body: UserIn, _: User = Depends(admin_only), db: Session = Depends(get_db)):
    if body.role not in ("admin", "collector"):
        raise HTTPException(400, "role admin ya collector hona chahiye")
    uname = body.username.strip().lower()
    if db.scalar(select(User).where(User.username == uname)):
        raise HTTPException(409, "Username already exists")
    u = User(username=uname, name=body.name.strip(), role=body.role, password_hash=hash_password(body.password))
    db.add(u)
    db.commit()
    return user_out(u)


@app.patch("/api/users/{uid}")
def patch_user(uid: int, body: UserPatch, admin: User = Depends(admin_only), db: Session = Depends(get_db)):
    u = db.get(User, uid)
    if not u:
        raise HTTPException(404, "User not found")
    data = body.model_dump(exclude_unset=True)
    if "role" in data and data["role"] not in ("admin", "collector"):
        raise HTTPException(400, "role admin ya collector hona chahiye")
    if u.id == admin.id and (data.get("active") is False or data.get("role") == "collector"):
        raise HTTPException(400, "Apna admin access khud band nahi kar sakte")
    if "name" in data and data["name"]:
        u.name = data["name"].strip()
    if "role" in data:
        u.role = data["role"]
    if "active" in data and data["active"] is not None:
        u.active = data["active"]
    if data.get("password"):
        u.password_hash = hash_password(data["password"])
    db.commit()
    return user_out(u)


@app.get("/api/collectors")
def collectors(_: User = Depends(current_user), db: Session = Depends(get_db)):
    return [{"id": u.id, "name": u.name} for u in db.scalars(select(User).where(User.active.is_(True)).order_by(User.name))]


# ------------------------------------------------------------------ sync
@app.get("/api/sync/status")
def sync_status(company: str = "", _: User = Depends(current_user), db: Session = Depends(get_db)):
    return svc.sync_status(db, resolve_company(company))


@app.post("/api/sync")
def sync_now(company: str = "", _: User = Depends(current_user), db: Session = Depends(get_db)):
    company = resolve_company(company)
    last = svc.meta_get(db, f"last_sync:{company}")
    if last and (config.now() - datetime.fromisoformat(last)).total_seconds() < 30:
        return {"synced": False, "reason": "recent", **svc.sync_status(db, company)}
    return svc.sync_customers(db, company, force=True)


# ------------------------------------------------------------------ customers
def _filtered(db, user, q, flt, assigned, status, company):
    ensure_fresh(db, company)
    customers = svc.company_customers(db, company)
    return svc.apply_filters(customers, q, flt, assigned, status, user.id)


@app.get("/api/customers")
def list_customers(
    q: str = "", flt: str = "with_balance", assigned: str = "", status: str = "",
    sort: str = "gold", desc: bool = True, limit: int = Query(1000, le=5000), company: str = "",
    user: User = Depends(current_user), db: Session = Depends(get_db),
):
    company = resolve_company(company)
    customers = _filtered(db, user, q, flt, assigned, status, company)
    users = svc.user_names(db)
    lasts, promises = svc.latest_followups(db), svc.pending_promises(db)
    prev = svc.prev_snapshot(db, company)
    rows = [svc.customer_dict(c, lasts.get(c.id), promises.get(c.id), users, prev=prev) for c in customers]
    rows = svc.sort_customers(rows, sort, desc)[:limit]
    return {"items": rows, "total": len(customers), "sync": svc.sync_status(db, company),
            "prev_day": prev[0].isoformat() if prev[0] else None}


@app.get("/api/customers/export.xlsx")
def export_customers(
    q: str = "", flt: str = "with_balance", assigned: str = "", status: str = "", company: str = "",
    user: User = Depends(current_user), db: Session = Depends(get_db),
):
    company = resolve_company(company)
    customers = _filtered(db, user, q, flt, assigned, status, company)
    users = svc.user_names(db)
    lasts = svc.latest_followups(db)
    wb = Workbook()
    ws = wb.active
    ws.title = "Customers"
    head = ["Code", "Name", "Name (AR)", "Mobile", "Address", "Gold balance (g)", "Cash balance", "Silver balance (g)",
            "Last invoice amt", "Last transaction", "Assigned to", "Status", "Last follow-up", "Next follow-up"]
    ws.append(head)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="5750F1")
    for c in sorted(customers, key=lambda c: c.gold_balance, reverse=True):
        lf = lasts.get(c.id)
        ws.append([c.code, c.name, c.name_ar, c.mobile or "", c.address or "", c.gold_balance, c.cash_balance, c.silver_balance,
                   c.last_invoice_amt, c.last_txn_date.isoformat() if c.last_txn_date else "",
                   users.get(c.assigned_to, ""), c.status,
                   lf.created_at.strftime("%Y-%m-%d") if lf else "",
                   lf.next_followup_date.isoformat() if lf and lf.next_followup_date else ""])
    for col, w in zip("ABCDEFGHIJKLMN", (14, 34, 34, 14, 30, 16, 16, 18, 16, 16, 18, 10, 16, 16)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    name = f"gold-customers-{company}-{config.today().isoformat()}.xlsx"
    return Response(buf.getvalue(), headers={"Content-Disposition": f'attachment; filename="{name}"'},
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


def _statement(cid, from_date, to_date, db):
    if from_date > to_date:
        raise HTTPException(400, "from_date to_date se pehle honi chahiye")
    if (to_date - from_date).days > 800:
        raise HTTPException(400, "Range bahut lambi hai (max ~2 saal)")
    c = db.get(Customer, cid)
    if not c:
        raise HTTPException(404, "Customer not found")
    co = svc.get_co(c.company)
    st = gerp.customer_statement(co, c.code, c.name, c.name_ar, from_date, to_date)
    st["company"] = co["name"]
    st["company_info"] = {"name": co["name"], "legal_name": co["legal_name"], "legal_name_en": co["legal_name_en"],
                          "vat": co["vat"], "address": co["address"], "address_en": co["address_en"]}
    st["generated_at"] = config.now().isoformat(timespec="minutes")
    return st


@app.get("/api/customers/{cid}/statement")
def customer_statement(cid: int, from_date: date, to_date: date, _: User = Depends(current_user),
                       db: Session = Depends(get_db)):
    return _statement(cid, from_date, to_date, db)


@app.get("/api/customers/{cid}/statement.xlsx")
def customer_statement_xlsx(cid: int, from_date: date, to_date: date, _: User = Depends(current_user),
                            db: Session = Depends(get_db)):
    st = _statement(cid, from_date, to_date, db)
    wb = Workbook()
    ws = wb.active
    ws.title = "Statement"
    ws.append(["Statement of account"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append(["Customer", st["customer"]["name"] or st["customer"]["name_ar"], st["customer"]["name_ar"] or ""])
    ws.append(["Code", st["customer"]["code"]])
    ws.append(["Period", f'{st["from_date"]} -> {st["to_date"]}'])
    ws.append(["Generated", st["generated_at"]])
    ws.append([])
    ws.append(["", "Opening", "Debit", "Credit", "Closing"])
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="5750F1")
    labels = {"sar": "Cash (SAR)", "c24": "Gold 24K (g)", "c22": "Gold 22K (g)", "c21": "Gold 21K (g)", "c18": "Gold 18K (g)",
              "gold24": "Gold total in 24K (g)", "silver": "Silver (g)"}
    for r in st["rows"]:
        ws.append([labels[r["key"]], r["opening"], r["debit"], r["credit"], r["closing"]])
    ws.append([])
    ws.append(["Balance = debit - credit. Positive = customer owes us. Summary statement (opening -> period -> closing); line-by-line entries are not available from the API."])
    for col, w in zip("ABCDE", (22, 18, 18, 18, 18)):
        ws.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    name = f'statement-{st["customer"]["code"] or cid}-{st["to_date"]}.xlsx'
    return Response(buf.getvalue(), headers={"Content-Disposition": f'attachment; filename="{name}"'},
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.post("/api/customers/import-report")
async def import_report(request: Request, company: str = "", force: bool = False, _: User = Depends(admin_only),
                        db: Session = Depends(get_db)):
    """Body = GERP 'Customer and Supplier Report' .xlsx (raw bytes). Admin only."""
    company = resolve_company(company)
    data = await request.body()
    if not data:
        raise HTTPException(400, "File khali hai")
    ensure_fresh(db, company)  # customers table bhari honi chahiye
    parsed = report_import.parse_report(data)
    res = report_import.apply_report(db, company, parsed, force=force)
    svc.meta_set(db, f"profile_import:{company}", config.now().isoformat(timespec="minutes"))
    db.commit()
    return {**res, "company": company}


def _ledger(cid, from_date, to_date, db):
    if from_date > to_date:
        raise HTTPException(400, "from_date to_date se pehle honi chahiye")
    if (to_date - from_date).days > 800:
        raise HTTPException(400, "Range bahut lambi hai (max ~2 saal)")
    c = db.get(Customer, cid)
    if not c:
        raise HTTPException(404, "Customer not found")
    co = svc.get_co(c.company)
    led = gerp.customer_ledger(co, c.code, from_date, to_date)
    led["customer"] = {"id": cid, "code": c.code, "name": c.name, "name_ar": c.name_ar}
    led["company"] = co["name"]
    led["company_info"] = {"name": co["name"], "legal_name": co["legal_name"], "legal_name_en": co["legal_name_en"],
                           "vat": co["vat"], "address": co["address"], "address_en": co["address_en"]}
    led["from_date"], led["to_date"] = from_date.isoformat(), to_date.isoformat()
    led["generated_at"] = config.now().isoformat(timespec="minutes")
    # last sync (customers table) se closing ka farq: nayi entry post hui ho to yahan dikhega
    led["vs_synced"] = {"cash": round(led["closing"]["cash"] - c.cash_balance, 4), "gold": round(led["closing"]["gold"] - c.gold_balance, 4),
                        "synced_at": c.synced_at.isoformat(timespec="minutes") if c.synced_at else None}
    return led


@app.get("/api/customers/{cid}/ledger")
def customer_ledger_view(cid: int, from_date: date, to_date: date, _: User = Depends(current_user), db: Session = Depends(get_db)):
    return _ledger(cid, from_date, to_date, db)


@app.get("/api/customers/{cid}/ledger.xlsx")
def customer_ledger_xlsx(cid: int, from_date: date, to_date: date, _: User = Depends(current_user), db: Session = Depends(get_db)):
    led = _ledger(cid, from_date, to_date, db)
    metals = led["active_metals"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Ledger"
    ws.append(["Statement of account - ledger"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append(["Customer", led["customer"]["name"] or led["customer"]["name_ar"], led["customer"]["name_ar"] or "", "Code", led["customer"]["code"]])
    ws.append(["Period", f'{led["from_date"]} -> {led["to_date"]}', "", "Company", led["company"]])
    ws.append([])
    head = ["#", "Date", "Voucher", "Type", "Type (AR)", "Remarks", "Debit", "Credit", "Balance (SAR)"]
    for k in metals:
        label = k.upper() if k in ("silver", "platinum", "palladium") else f"{k}K"
        head += [f"{label} Debit", f"{label} Credit", f"{label} Balance"]
    ws.append(head)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="5750F1")
    for l in led["lines"]:
        row = [l["index"], l["date"], l["voucher_no"], l["type"], l["type_ar"], l["remarks"], l["debit"], l["credit"], l["balance"]]
        for k in metals:
            row += [l["metals"][k]["debit"], l["metals"][k]["credit"], l["metals"][k]["balance"]]
        ws.append(row)
    tot = ["", "", "", "", "", "Total", led["totals"]["cash"]["debit"], led["totals"]["cash"]["credit"], led["totals"]["cash"]["balance"]]
    for k in metals:
        tot += [led["totals"][k]["debit"], led["totals"][k]["credit"], led["totals"][k]["balance"]]
    ws.append(tot)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
    ws.freeze_panes = ws.cell(row=6, column=1)
    for i, w in enumerate([5, 12, 16, 22, 18, 40, 16, 16, 16] + [14] * (3 * len(metals)), start=1):
        ws.column_dimensions[ws.cell(row=5, column=i).column_letter].width = w
    buf = io.BytesIO()
    wb.save(buf)
    name = f'ledger-{led["customer"]["code"]}-{led["to_date"]}.xlsx'
    return Response(buf.getvalue(), headers={"Content-Disposition": f'attachment; filename="{name}"'},
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


class AssignIn(BaseModel):
    ids: List[int]
    user_id: Optional[int] = None


@app.post("/api/customers/assign")
def bulk_assign(body: AssignIn, _: User = Depends(admin_only), db: Session = Depends(get_db)):
    if body.user_id is not None and not db.get(User, body.user_id):
        raise HTTPException(404, "User not found")
    n = 0
    for c in db.scalars(select(Customer).where(Customer.id.in_(body.ids))):
        c.assigned_to = body.user_id
        n += 1
    db.commit()
    return {"updated": n}


@app.get("/api/customers/{cid}")
def customer_detail(cid: int, user: User = Depends(current_user), db: Session = Depends(get_db)):
    c = db.get(Customer, cid)
    if not c:
        raise HTTPException(404, "Customer not found")
    users = svc.user_names(db)
    fus = db.scalars(select(Followup).where(Followup.customer_id == cid).order_by(Followup.id.desc())).all()
    last = fus[0] if fus else None
    promise = next((f for f in sorted(fus, key=lambda f: f.promise_date or date.max)
                    if f.promise_status == "pending" and f.promise_date), None)
    return {**svc.customer_dict(c, last, promise, users, detail=True, prev=svc.prev_snapshot(db, c.company)),
            "followups": [svc.followup_dict(f) for f in fus],
            "history": svc.customer_history(db, cid)}


class CustomerPatch(BaseModel):
    assigned_to: Optional[int] = None
    status: Optional[str] = None
    notes: Optional[str] = Field(default=None, max_length=4000)


@app.patch("/api/customers/{cid}")
def patch_customer(cid: int, body: CustomerPatch, user: User = Depends(current_user), db: Session = Depends(get_db)):
    c = db.get(Customer, cid)
    if not c:
        raise HTTPException(404, "Customer not found")
    data = body.model_dump(exclude_unset=True)
    if "assigned_to" in data:
        target = data["assigned_to"]
        if user.role != "admin" and target not in (None, user.id):
            raise HTTPException(403, "Collector sirf apne aap ko assign kar sakta hai")
        if user.role != "admin" and target is None and c.assigned_to not in (None, user.id):
            raise HTTPException(403, "Kisi aur ka assignment sirf admin hata sakta hai")
        if target is not None and not db.get(User, target):
            raise HTTPException(404, "User not found")
        c.assigned_to = target
    if "status" in data:
        if data["status"] not in STATUSES:
            raise HTTPException(400, "status invalid")
        c.status = data["status"]
    if "notes" in data and data["notes"] is not None:
        c.notes = data["notes"]
    db.commit()
    return {"ok": True}


# ------------------------------------------------------------------ followups & promises
class FollowupIn(BaseModel):
    kind: str = "call"
    outcome: str = "other"
    note: str = Field(default="", max_length=4000)
    next_followup_date: Optional[date] = None
    promise_date: Optional[date] = None
    promise_gold: Optional[float] = Field(default=None, ge=0)
    promise_cash: Optional[float] = Field(default=None, ge=0)
    promise_silver: Optional[float] = Field(default=None, ge=0)


@app.post("/api/customers/{cid}/followups")
def add_followup(cid: int, body: FollowupIn, user: User = Depends(current_user), db: Session = Depends(get_db)):
    if not db.get(Customer, cid):
        raise HTTPException(404, "Customer not found")
    if body.kind not in KINDS or body.outcome not in OUTCOMES:
        raise HTTPException(400, "kind/outcome invalid")
    has_amount = any(v for v in (body.promise_gold, body.promise_cash, body.promise_silver))
    if body.promise_date is None and has_amount:
        raise HTTPException(400, "Promise ke saath date zaroori hai")
    if body.promise_date and not has_amount:
        raise HTTPException(400, "Promise mein gold/cash/silver ka amount daalo")
    f = Followup(
        customer_id=cid, user_id=user.id, kind=body.kind, outcome=body.outcome, note=body.note.strip(),
        next_followup_date=body.next_followup_date, promise_date=body.promise_date,
        promise_gold=body.promise_gold or None, promise_cash=body.promise_cash or None,
        promise_silver=body.promise_silver or None,
        promise_status="pending" if body.promise_date else None,
    )
    db.add(f)
    db.commit()
    db.refresh(f)
    return svc.followup_dict(f)


class FollowupPatch(BaseModel):
    promise_status: Optional[str] = None
    note: Optional[str] = Field(default=None, max_length=4000)
    next_followup_date: Optional[date] = None


@app.patch("/api/followups/{fid}")
def patch_followup(fid: int, body: FollowupPatch, user: User = Depends(current_user), db: Session = Depends(get_db)):
    f = db.get(Followup, fid)
    if not f:
        raise HTTPException(404, "Follow-up not found")
    data = body.model_dump(exclude_unset=True)
    if "promise_status" in data:
        if not f.promise_date or data["promise_status"] not in PROMISE_STATUSES:
            raise HTTPException(400, "promise_status invalid")
        f.promise_status = data["promise_status"]
    if "note" in data and data["note"] is not None:
        if user.role != "admin" and f.user_id != user.id:
            raise HTTPException(403, "Sirf author ya admin note edit kar sakta hai")
        f.note = data["note"]
    if "next_followup_date" in data:
        f.next_followup_date = data["next_followup_date"]
    db.commit()
    return svc.followup_dict(f)


@app.delete("/api/followups/{fid}")
def delete_followup(fid: int, user: User = Depends(current_user), db: Session = Depends(get_db)):
    f = db.get(Followup, fid)
    if not f:
        raise HTTPException(404, "Follow-up not found")
    if user.role != "admin" and f.user_id != user.id:
        raise HTTPException(403, "Sirf author ya admin delete kar sakta hai")
    db.delete(f)
    db.commit()
    return {"ok": True}


# ------------------------------------------------------------------ dashboards
@app.get("/api/myday")
def my_day(scope: str = "", company: str = "", user: User = Depends(current_user), db: Session = Depends(get_db)):
    company = resolve_company(company)
    ensure_fresh(db, company)
    scope = scope if scope in ("mine", "all") else ("all" if user.role == "admin" else "mine")
    return svc.my_day(db, user, scope, company)


@app.get("/api/overview")
def overview(company: str = "", _: User = Depends(current_user), db: Session = Depends(get_db)):
    company = resolve_company(company)
    ensure_fresh(db, company)
    return {**svc.overview(db, company), "sync": svc.sync_status(db, company)}


@app.get("/api/kpi")
def kpi(
    from_date: str = Query(..., pattern=r"^\d{4}-\d{2}-\d{2}$"),
    to_date: str = Query(..., pattern=r"^\d{4}-\d{2}-\d{2}$"),
    branches: str = "", company: str = "",
    opening_gold_grams: float = 0, opening_gold_cash: float = 0, opening_silver: float = 0,
    _: User = Depends(current_user),
):
    try:
        co = svc.get_co(company)
        ids = [int(b) for b in branches.split(",") if b.strip()] or co["branches"]
    except ValueError:
        raise HTTPException(400, "branches sirf numbers hone chahiye")
    return gerp.kpi(co, ids, from_date, to_date, opening_gold_grams, opening_gold_cash, opening_silver)
